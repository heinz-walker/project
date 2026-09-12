#!/usr/bin/env python3
import argparse
import csv
import datetime as dt
import gzip
import html
import ipaddress
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from string import Template
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SEOUL = ZoneInfo("Asia/Seoul")
SSM_SUCCESS_EVENTS = {"StartSession", "ResumeSession", "TerminateSession"}
ANSI_RE = re.compile(r"\x1B\[[0-?]*[ -/]*[@-~]|\x1B\][^\x07]*\x07|\x1b=")
CONTROL_RE = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")
PROMPT_RE = re.compile(r"^(?:\[[^\]]+\][#$]|mysql>|sh-\d+(?:\.\d+)?\$|[A-Za-z0-9_.-]+@[^ ]+[#$])\s*")
RISK_PATTERNS = [
    ("credential_exposure", re.compile(r"password=|passwd|secret|token|private[_-]?key", re.I)),
    ("db_access", re.compile(r"\b(mysql|psql|redis-cli|mongosh?)\b", re.I)),
    ("privilege_change", re.compile(r"\b(sudo|su\s|chmod|chown|usermod|visudo)\b", re.I)),
    ("service_control", re.compile(r"\b(systemctl|service|docker|kubectl)\b", re.I)),
    ("data_archive", re.compile(r"\b(tar|zip|gzip|scp|rsync|mysqldump|pg_dump)\b", re.I)),
    ("destructive", re.compile(r"\b(rm\s+-rf|truncate|drop\s+table|delete\s+from)\b", re.I)),
]
RISK_PATTERN_LABELS = {
    "credential_exposure": "자격증명 노출",
    "db_access": "DB 직접 접근",
    "privilege_change": "권한 변경",
    "service_control": "서비스 제어",
    "data_archive": "데이터 아카이브",
    "destructive": "파괴적 명령",
}
KNOWN_START_COMMANDS = (
    "cd ~ && exec /bin/bash",
    "bash",
    "sh",
    "/bin/bash",
)
KNOWN_COMMAND_PREFIXES = (
    "cd ", "ls", "cat ", "alias", "./", "mysql", "psql", "redis-cli",
    "grep ", "find ", "pwd", "sudo ", "su ", "ssh ", "scp ", "rsync ",
    "tar ", "zip ", "rm ", "mv ", "cp ", "kubectl ", "docker ", "systemctl ",
)
WEBAPP_SHEETS = {
    "webapp-a-admin": "게임 서비스 A",
    "webapp-b-admin": "게임 서비스 B",
    "webapp-c-admin": "게임 서비스 C",
}
SENSITIVE_WEB_PATTERNS = [
    ("refund_state_change", re.compile(r"상태 변경", re.I)),
    ("asset_change", re.compile(r"재화 변경", re.I)),
    ("vip_change", re.compile(r"VIP 변경", re.I)),
    ("memo_add", re.compile(r"메모 추가", re.I)),
    ("suspension", re.compile(r"정지 상태 처리", re.I)),
    ("item_change", re.compile(r"아이템 변경", re.I)),
    ("message_send", re.compile(r"메시지발송", re.I)),
]


def run_aws(args):
    proc = subprocess.run(["aws", *args], capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or "aws command failed")
    return proc.stdout


def run_aws_json(args):
    return json.loads(run_aws(args))


def daterange(start_date, end_date):
    current = start_date
    while current <= end_date:
        yield current
        current += dt.timedelta(days=1)


def sync_day_logs(bucket, account_id, region, day, profile, local_dir):
    prefix = f"AWSLogs/{account_id}/CloudTrail/{region}/{day:%Y/%m/%d}/"
    if local_dir.exists():
        shutil.rmtree(local_dir)
    local_dir.mkdir(parents=True, exist_ok=True)
    run_aws([
        "s3",
        "sync",
        f"s3://{bucket}/{prefix}",
        str(local_dir),
        "--exclude",
        "*",
        "--include",
        "*.json.gz",
        "--profile",
        profile,
    ])


def load_cloudtrail_records(local_file):
    payload = gzip.decompress(local_file.read_bytes())
    return json.loads(payload.decode("utf-8")).get("Records", [])


def to_kst(timestamp):
    return dt.datetime.fromisoformat(timestamp.replace("Z", "+00:00")).astimezone(SEOUL)


def format_kst_label(moment):
    return moment.strftime("%Y년 %m월 %d일 %H:%M:%S")


def is_after_hours(moment, start_hour, end_hour):
    if moment.weekday() >= 5:
        return True
    start = moment.replace(hour=start_hour, minute=0, second=0, microsecond=0)
    end = moment.replace(hour=end_hour, minute=0, second=0, microsecond=0)
    return moment < start or moment > end


def load_config(config_path):
    if not config_path or not Path(config_path).exists():
        return {}
    return json.loads(Path(config_path).read_text(encoding="utf-8"))


def normalize_user_name(raw_user_name, config):
    if not raw_user_name:
        return ""
    aliases = config.get("user_aliases") or {}
    if raw_user_name in aliases:
        return aliases[raw_user_name]
    if ":" in raw_user_name:
        suffix = raw_user_name.rsplit(":", 1)[-1]
        return aliases.get(suffix, suffix)
    return raw_user_name


def is_holiday(moment, config):
    holidays = set(config.get("holidays") or [])
    return moment.date().isoformat() in holidays


def is_ip_allowed(source_ip, config):
    if not source_ip:
        return False
    if source_ip in (config.get("allowed_ips") or []):
        return True
    for cidr in config.get("allowed_ip_cidrs") or []:
        try:
            if ipaddress.ip_address(source_ip) in ipaddress.ip_network(cidr, strict=False):
                return True
        except ValueError:
            continue
    return False


def normalize_admin_label(label):
    if not label:
        return ""
    text = str(label)
    match = re.match(r"^(.*?)\s+\((.*?)\)$", text)
    if match:
        return {
            "display_name": match.group(1).strip(),
            "user_id": match.group(2).strip(),
            "full": text.strip(),
        }
    return {
        "display_name": text.strip(),
        "user_id": text.strip(),
        "full": text.strip(),
    }


def match_rule(rule, row):
    for key, expected in rule.items():
        if key in {"status", "reason", "remark"}:
            continue
        actual = str(row.get(key, ""))
        if key.endswith("_regex"):
            base_key = key[:-6]
            actual = str(row.get(base_key, ""))
            if not re.search(expected, actual):
                return False
            continue
        if actual != str(expected):
            return False
    return True


def find_matching_rule(rules, row):
    for rule in rules or []:
        if match_rule(rule, row):
            return rule
    return None


def extract_cloudtrail_findings(records, business_start_hour, business_end_hour, config):
    success_rows = []
    blocked_rows = []
    for record in records:
        if record.get("eventSource") != "ssm.amazonaws.com":
            continue
        event_name = record.get("eventName", "")
        if event_name not in SSM_SUCCESS_EVENTS:
            continue

        user_identity = record.get("userIdentity", {}) or {}
        request_params = record.get("requestParameters", {}) or {}
        response_elements = record.get("responseElements", {}) or {}
        kst_time = to_kst(record["eventTime"])
        raw_user_name = user_identity.get("userName") or user_identity.get("principalId", "")
        user_name = normalize_user_name(raw_user_name, config)
        holiday = is_holiday(kst_time, config)

        base = {
            "event_time_kst": kst_time.isoformat(timespec="seconds"),
            "event_time_display": format_kst_label(kst_time),
            "event_date_kst": kst_time.date().isoformat(),
            "event_name": event_name,
            "user_name": user_name,
            "raw_user_name": raw_user_name,
            "user_arn": user_identity.get("arn", ""),
            "source_ip": record.get("sourceIPAddress", ""),
            "target_instance_id": request_params.get("target", ""),
            "session_id": response_elements.get("sessionId", ""),
            "mfa_authenticated": user_identity.get("sessionContext", {}).get("attributes", {}).get("mfaAuthenticated", ""),
            "aws_region": record.get("awsRegion", ""),
            "after_hours": "Y" if holiday or is_after_hours(kst_time, business_start_hour, business_end_hour) else "N",
            "weekend": "Y" if kst_time.weekday() >= 5 else "N",
            "holiday": "Y" if holiday else "N",
            "source_ip_allowed": "Y" if is_ip_allowed(record.get("sourceIPAddress", ""), config) else "N",
            "user_authorized": "Y" if user_name in (config.get("authorized_users") or {}) else "N",
            "event_id": record.get("eventID", ""),
        }

        if record.get("errorCode"):
            blocked = dict(base)
            blocked["error_code"] = record.get("errorCode", "")
            blocked["error_message"] = record.get("errorMessage", "")
            blocked_rows.append(blocked)
            continue

        success_rows.append(base)

    success_rows.sort(key=lambda row: row["event_time_kst"])
    blocked_rows.sort(key=lambda row: row["event_time_kst"])
    return success_rows, blocked_rows


def clean_text(text):
    text = ANSI_RE.sub("", text)
    text = CONTROL_RE.sub("", text)
    return text.strip()


def normalize_command(line):
    line = clean_text(line)
    while True:
        updated = PROMPT_RE.sub("", line)
        if updated == line:
            break
        line = updated.strip()
    return line.strip()


def classify_command(command):
    hits = [name for name, pattern in RISK_PATTERNS if pattern.search(command)]
    return ",".join(hits)


def get_stream_events(profile, log_group, stream_name):
    try:
        result = run_aws_json([
            "logs",
            "filter-log-events",
            "--profile",
            profile,
            "--log-group-name",
            log_group,
            "--log-stream-names",
            stream_name,
        ])
    except RuntimeError as exc:
        if "ResourceNotFoundException" in str(exc):
            return []
        raise
    return result.get("events", [])


def extract_command_rows(profile, log_group, success_rows):
    commands = []
    seen = set()
    for session in success_rows:
        session_id = session.get("session_id")
        if not session_id:
            continue
        for event in get_stream_events(profile, log_group, session_id):
            message = json.loads(event["message"])
            session_data = message.get("sessionData", []) or []
            if not session_data:
                continue
            raw_line = session_data[0]
            line = clean_text(raw_line)
            if not line:
                continue
            command = normalize_command(line)
            if not command:
                continue
            lower = command.lower()
            if lower.startswith(("welcome to ", "last login:", "system information as of ", "usage of /:", "memory usage:", "swap usage:")):
                continue
            if command in {"$", "#", "mysql>"} or command.endswith("$") or command.endswith("#"):
                continue
            if not (command.startswith(KNOWN_START_COMMANDS) or lower.startswith(KNOWN_COMMAND_PREFIXES)):
                continue
            row = {
                "event_time": message.get("eventTime", ""),
                "event_time_kst": format_kst_label(to_kst(message.get("eventTime", ""))),
                "session_id": message.get("sessionId", ""),
                "user_arn": message.get("userIdentity", {}).get("arn", ""),
                "user_name": session.get("user_name", ""),
                "run_as_user": message.get("runAsUser", ""),
                "target_instance_id": message.get("target", {}).get("id", ""),
                "aws_region": message.get("awsRegion", ""),
                "command": command,
                "risk_flags": classify_command(command),
            }
            key = (row["session_id"], row["event_time"], row["command"])
            if key in seen:
                continue
            seen.add(key)
            commands.append(row)
    commands.sort(key=lambda row: (row["event_time"], row["session_id"], row["command"]))
    return commands


def classify_webapp_sensitive(content):
    for label, pattern in SENSITIVE_WEB_PATTERNS:
        if pattern.search(content):
            return label
    return ""


def extract_webapp_findings(workbook_path, start_date, end_date, business_start_hour, business_end_hour, config):
    if not workbook_path or not Path(workbook_path).exists():
        return None

    wb = load_workbook(workbook_path, data_only=True)
    service_summary = []
    after_hours_rows = []
    sensitive_rows = []
    unauthorized_webapp_rows = []
    service_unresolved = {}
    webapp_auth = config.get("webapp_authorized_users") or {}

    for sheet_name, service_name in WEBAPP_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        login_count = 0
        after_hours_count = 0
        sensitive_count = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 6:
                continue
            _, event_time, admin, source_ip, target_user, content = row[:6]
            if not event_time or not content:
                continue
            if not isinstance(event_time, dt.datetime):
                continue
            if not (start_date <= event_time.date() <= end_date):
                continue
            admin_info = normalize_admin_label(admin)
            display_time = format_kst_label(event_time.replace(tzinfo=SEOUL) if event_time.tzinfo is None else event_time.astimezone(SEOUL))
            content_text = str(content)

            if content_text == "로그인":
                login_count += 1
                if webapp_auth and admin_info["user_id"] not in webapp_auth:
                    unauthorized_webapp_rows.append({
                        "service": service_name,
                        "event_time_display": display_time,
                        "admin_label": admin_info["full"],
                        "source_ip": str(source_ip),
                        "review_status": "비인가 의심",
                        "review_summary": "인가자 목록에 없는 관리자 로그인",
                        "remark": "접속 목적 확인 필요",
                    })
                holiday = is_holiday(event_time.replace(tzinfo=SEOUL), config)
                is_outside = holiday or is_after_hours(event_time.replace(tzinfo=SEOUL), business_start_hour, business_end_hour)
                if is_outside:
                    after_hours_count += 1
                    candidate = {
                        "service": service_name,
                        "admin_label": admin_info["full"],
                        "source_ip": str(source_ip),
                    }
                    auto_rule = find_matching_rule(config.get("webapp_after_hours_auto_rules"), candidate)
                    if auto_rule:
                        review_status = auto_rule.get("status", "확인 필요")
                        review_summary = auto_rule.get("reason", "웹앱 업무시간 외 로그인 이력")
                        remark = auto_rule.get("remark", "")
                    else:
                        review_status = "확인 필요"
                        review_summary = "웹앱 업무시간 외 로그인 이력"
                        remark = "업무 목적 또는 승인 근거 확인 필요"
                    if review_status != "이상없음(양호)":
                        service_unresolved[service_name] = True
                    elif service_name not in service_unresolved:
                        service_unresolved[service_name] = False
                    after_hours_rows.append({
                        "service": service_name,
                        "event_time_display": display_time,
                        "admin_label": admin_info["full"],
                        "source_ip": str(source_ip),
                        "review_status": review_status,
                        "review_summary": review_summary,
                        "remark": remark,
                    })

            sensitive_flag = classify_webapp_sensitive(content_text)
            if sensitive_flag:
                sensitive_count += 1
                sensitive_rows.append({
                    "service": service_name,
                    "event_time_display": display_time,
                    "admin_label": admin_info["full"],
                    "source_ip": str(source_ip),
                    "review_status": "이상없음(양호)",
                    "review_summary": f"{sensitive_flag} 처리행위 로그 기록 확인",
                    "remark": "상세 행위 내역은 첨부 XLSX 참조",
                })

        has_unresolved = service_unresolved.get(service_name, False)
        service_summary.append({
            "service": service_name,
            "login_count": login_count,
            "after_hours_count": after_hours_count,
            "sensitive_count": sensitive_count,
            "review_status": "확인 필요" if has_unresolved else "이상없음(양호)",
            "review_summary": (
                "업무시간 외 로그인 이력 존재 (미해소)"
                if has_unresolved
                else (
                    "업무시간 외 로그인이 모두 업무 목적 확인 완료됨"
                    if after_hours_count > 0
                    else "로그인 및 주요 처리행위 로그 범위 내 특이사항 없음"
                )
            ),
            "remark": "실패 로그인/차단 로그는 본 엑셀 원본에 없음",
        })

    web_role_change_rows = []
    if "게임Admin 권한 변경 로그" in wb.sheetnames:
        ws = wb["게임Admin 권한 변경 로그"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 5:
                continue
            event_time, target, message, admin, game = row[:5]
            if not isinstance(event_time, dt.datetime):
                continue
            if not (start_date <= event_time.date() <= end_date):
                continue
            web_role_change_rows.append({
                "service": game or "게임Admin 권한 변경",
                "event_time_display": format_kst_label(event_time.replace(tzinfo=SEOUL)),
                "admin_label": str(admin or "-"),
                "source_ip": "-",
                "review_status": "확인 필요",
                "review_summary": f"권한 변경 로그 확인: {message}",
                "remark": f"대상: {target}",
            })

    limit_rows = [
        *([] if webapp_auth else [{"item": "비인가자 접속 성공 유무", "content": "원본 엑셀에는 웹앱 인가자 기준 정보가 없으므로 별도 인가자 목록이 필요함"}]),
        {"item": "불필요 접속시도 차단 현황", "content": "원본 엑셀에는 실패 로그인 또는 접근 차단 이벤트가 포함되어 있지 않음"},
    ]

    after_hours_rows.sort(key=lambda row: row["event_time_display"])
    sensitive_rows.sort(key=lambda row: row["event_time_display"], reverse=True)
    web_role_change_rows.sort(key=lambda row: row["event_time_display"], reverse=True)

    return {
        "service_summary": service_summary,
        "unauthorized_webapp_rows": unauthorized_webapp_rows,
        "after_hours_rows": after_hours_rows,
        "sensitive_rows": sensitive_rows[:20],
        "role_change_rows": web_role_change_rows[:20],
        "limit_rows": limit_rows,
    }


def write_csv(rows, out_path, fieldnames):
    with out_path.open("w", newline="", encoding="utf-8-sig") as fp:
        writer = csv.DictWriter(fp, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_json(data, out_path):
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_template(template_path):
    return Template(Path(template_path).read_text(encoding="utf-8"))


TD_STYLE = "border:1px solid #222;padding:8px 10px;vertical-align:top;word-break:break-word;"
STATUS_COLORS = {"good": "#1f5fae", "warn": "#c45500", "bad": "#b42318"}


def html_list(items):
    lis = "".join(f'<li style="margin:2px 0">{item}</li>' for item in items)
    return f'<ul style="margin:6px 0 12px 18px;padding:0">{lis}</ul>'


def html_text(value):
    return html.escape(str(value)) if value not in (None, "") else "-"


def html_status(value):
    css = "good"
    if "확인" in value:
        css = "warn"
    if "비인가" in value or "위험" in value:
        css = "bad"
    color = STATUS_COLORS[css]
    return f'<span style="color:{color};font-weight:700">{html.escape(value)}</span>'


def make_row(cells):
    return "<tr>" + "".join(cells) + "</tr>"


def td(value, cls=""):
    style = TD_STYLE
    if "center" in cls:
        style += "text-align:center;"
    if "muted" in cls:
        style += "color:#666;"
    return f'<td style="{style}">{value}</td>'


def no_data_row(message, colspan=6):
    return f'<tr><td colspan="{colspan}" style="{TD_STYLE}text-align:center;color:#666;">{html.escape(message)}</td></tr>'


def authorized_user_state(config):
    return bool(config.get("authorized_users"))


def get_user_meta(config, user_name):
    return (config.get("authorized_users") or {}).get(user_name, {})


def ip_label(config, source_ip):
    label = (config.get("ip_labels") or {}).get(source_ip)
    return f"{source_ip} ({label})" if label else source_ip


def session_approval(config, section, session_id):
    return ((config.get(section) or {}).get(session_id)) or {}


def summarize(success_rows, blocked_rows, risky_rows, config):
    unauthorized_findings = []
    after_hours_findings = []
    risky_findings = []

    if authorized_user_state(config):
        for row in success_rows:
            if row["user_name"] not in config["authorized_users"]:
                unauthorized_findings.append(row)

    for row in success_rows:
        if row["after_hours"] == "Y":
            approval = session_approval(config, "after_hours_approvals", row["session_id"])
            auto_rule = find_matching_rule(config.get("after_hours_auto_rules"), row)
            applied = approval or auto_rule or {}
            status = applied.get("status", "이상없음(양호)" if applied else "확인 필요")
            enriched = dict(row)
            enriched["review_status"] = status
            enriched["review_note"] = applied.get("reason", "업무 목적/사전 승인 근거 확인 필요")
            enriched["remark"] = applied.get("remark", "")
            after_hours_findings.append(enriched)

    for row in risky_rows:
        if not row["risk_flags"]:
            continue
        approval = session_approval(config, "command_approvals", row["session_id"])
        auto_rule = find_matching_rule(config.get("command_auto_rules"), row)
        applied = approval or auto_rule or {}
        status = applied.get("status", "이상없음(양호)" if applied else "확인 필요")
        enriched = dict(row)
        enriched["review_status"] = status
        enriched["review_note"] = applied.get("reason", f"위험 패턴({row['risk_flags']})에 대한 업무 필요성 확인 필요")
        enriched["remark"] = applied.get("remark", "")
        if enriched["review_status"] != "제외":
            risky_findings.append(enriched)

    blocked_findings = []
    for row in blocked_rows:
        approval = session_approval(config, "blocked_approvals", row["event_id"])
        auto_rule = find_matching_rule(config.get("blocked_auto_rules"), row)
        applied = approval or auto_rule or {}
        enriched = dict(row)
        enriched["review_status"] = applied.get("status", "확인 필요")
        enriched["review_note"] = applied.get("reason", row.get("error_code", "차단 사유 확인 필요"))
        enriched["remark"] = applied.get("remark", "")
        if enriched["review_status"] != "제외":
            blocked_findings.append(enriched)

    summary_rows = [
        {
            "name": "비인가자 접속 성공 유무",
            "count": len(unauthorized_findings) if authorized_user_state(config) else len(success_rows),
            "status": "이상없음(양호)" if authorized_user_state(config) and not unauthorized_findings else ("검토기준 미설정" if not authorized_user_state(config) else "확인 필요"),
            "note": "인가자 기준 파일 미설정" if not authorized_user_state(config) else ("비인가 접속 성공 이력 없음" if not unauthorized_findings else "인가자 목록에 없는 사용자 접속 확인"),
        },
        {
            "name": "업무시간 외 접속이력",
            "count": len(after_hours_findings),
            "status": "이상없음(양호)" if not after_hours_findings or all(item["review_status"] == "이상없음(양호)" for item in after_hours_findings) else "확인 필요",
            "note": (
                "업무시간 외 접속 없음"
                if not after_hours_findings
                else (
                    "업무시간 외 접속이 모두 승인 처리됨"
                    if all(item["review_status"] == "이상없음(양호)" for item in after_hours_findings)
                    else "업무시간 외 접속 건별 근거 확인 필요"
                )
            ),
        },
        {
            "name": "과도한 명령어 수행",
            "count": len(risky_findings),
            "status": "이상없음(양호)" if not risky_findings or all(item["review_status"] == "이상없음(양호)" for item in risky_findings) else "확인 필요",
            "note": (
                "위험 패턴 명령 미탐지"
                if not risky_findings
                else (
                    "위험 패턴 명령이 모두 승인 규칙 또는 예외 근거로 처리됨"
                    if all(item["review_status"] == "이상없음(양호)" for item in risky_findings)
                    else "위험 패턴 명령 건별 업무 필요성 확인 필요"
                )
            ),
        },
        {
            "name": "불필요 접속시도 차단 현황",
            "count": len(blocked_findings),
            "status": "이상없음(양호)" if not blocked_findings or all(item["review_status"] == "이상없음(양호)" for item in blocked_findings) else "확인 필요",
            "note": (
                "차단된 StartSession 이력 없음"
                if not blocked_findings
                else (
                    "차단 이력이 모두 정상 차단 또는 설정 검증성 실패로 분류됨"
                    if all(item["review_status"] == "이상없음(양호)" for item in blocked_findings)
                    else "차단 이력의 사용자/사유 확인 필요"
                )
            ),
        },
    ]

    return summary_rows, unauthorized_findings, after_hours_findings, risky_findings, blocked_findings


def render_summary_rows(summary_rows):
    rows = []
    for row in summary_rows:
        rows.append(make_row([
            td(html.escape(row["name"])),
            td(html.escape(str(row["count"])), "center"),
            td(html_status(row["status"]), "center"),
            td(html.escape(row["note"])),
        ]))
    return "".join(rows)


def render_webapp_summary_rows(rows):
    if not rows:
        return no_data_row("웹앱 로그 원본이 없어 검토 요약을 생성하지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        out.append(make_row([
            td(html_text(row["service"])),
            td(html_text(row["login_count"]), "center"),
            td(html_text(row["after_hours_count"]), "center"),
            td(html_text(row["sensitive_count"]), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(row["review_summary"])),
            td(html_text(row["remark"])),
        ]))
    return "".join(out)


def render_webapp_event_rows(rows, empty_message):
    if not rows:
        return no_data_row(empty_message, colspan=7)
    out = []
    for row in rows:
        out.append(make_row([
            td(html_text(row["service"])),
            td(html_text(row["event_time_display"]), "center"),
            td(html_text(row["admin_label"]), "center"),
            td(html_text(row["source_ip"]), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(row["review_summary"])),
            td(html_text(row["remark"])),
        ]))
    return "".join(out)


def render_webapp_unauthorized_rows(rows, has_auth_config):
    if not has_auth_config:
        return no_data_row("인가자 기준 파일이 없어 자동 판정을 수행하지 않았습니다.", colspan=7)
    if not rows:
        return no_data_row("웹앱 비인가자 접속 성공 이력이 확인되지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        out.append(make_row([
            td(html_text(row["service"])),
            td(html_text(row["event_time_display"]), "center"),
            td(html_text(row["admin_label"]), "center"),
            td(html_text(row["source_ip"]), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(row["review_summary"])),
            td(html_text(row["remark"])),
        ]))
    return "".join(out)


def render_webapp_limit_rows(rows):
    if not rows:
        return no_data_row("검토 한계 없음", colspan=2)
    out = []
    for row in rows:
        out.append(make_row([
            td(html_text(row["item"])),
            td(html_text(row["content"])),
        ]))
    return "".join(out)


def render_unauthorized_rows(rows, config):
    if not authorized_user_state(config):
        return no_data_row("인가자 기준 파일이 없어 자동 판정을 수행하지 않았습니다.", colspan=7)
    if not rows:
        return no_data_row("비인가자 접속 성공 이력이 확인되지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        out.append(make_row([
            td("PIMS bastion host"),
            td(html_text(row["event_time_display"]), "center"),
            td(html_text(row["user_name"]), "center"),
            td(html_text(ip_label(config, row["source_ip"])), "center"),
            td(html_status("비인가 의심"), "center"),
            td("인가자 목록에 없는 사용자로 확인되어 검토 필요"),
            td("-"),
        ]))
    return "".join(out)


def render_after_hours_rows(rows, config):
    if not rows:
        return no_data_row("업무시간 외 접속 이력이 확인되지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        out.append(make_row([
            td("PIMS bastion host"),
            td(html_text(row["event_time_display"]), "center"),
            td(html_text(row["user_name"]), "center"),
            td(html_text(ip_label(config, row["source_ip"])), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(row["review_note"])),
            td(html_text(row.get("remark", ""))),
        ]))
    return "".join(out)


def render_risky_rows(rows):
    if not rows:
        return no_data_row("위험 패턴으로 분류된 명령 이력이 확인되지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        summary = row["review_note"]
        if row.get("risk_flags"):
            summary = f"{row['risk_flags']} 패턴 명령"
        out.append(make_row([
            td("PIMS bastion host"),
            td(html_text(row["event_time_kst"]), "center"),
            td(html_text(row["user_arn"].rsplit("/", 1)[-1]), "center"),
            td(html_text(row["session_id"]), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(summary)),
            td(html_text(row.get("remark", ""))),
        ]))
    return "".join(out)


def render_risky_summary_rows(risky_findings):
    if not risky_findings:
        return no_data_row("위험 패턴으로 분류된 명령 이력이 확인되지 않았습니다.", colspan=4)
    from collections import defaultdict
    counts = defaultdict(lambda: {"total": 0, "unresolved": 0})
    for row in risky_findings:
        for flag in (f.strip() for f in row["risk_flags"].split(",") if f.strip()):
            counts[flag]["total"] += 1
            if row["review_status"] != "이상없음(양호)":
                counts[flag]["unresolved"] += 1
    out = []
    for flag, _ in RISK_PATTERNS:
        if flag not in counts:
            continue
        c = counts[flag]
        status = "확인 필요" if c["unresolved"] > 0 else "이상없음(양호)"
        remark = f"미해소 {c['unresolved']}건 확인 필요" if c["unresolved"] > 0 else "상세 내역 XLSX 참조"
        label = RISK_PATTERN_LABELS.get(flag, flag)
        out.append(make_row([
            td(html_text(label)),
            td(html_text(c["total"]), "center"),
            td(html_status(status), "center"),
            td(html_text(remark)),
        ]))
    return "".join(out)


def render_blocked_rows(rows, config):
    if not rows:
        return no_data_row("차단된 StartSession 이력이 확인되지 않았습니다.", colspan=7)
    out = []
    for row in rows:
        note = row["review_note"] or row["error_code"]
        out.append(make_row([
            td("PIMS bastion host"),
            td(html_text(row["event_time_display"]), "center"),
            td(html_text(row["user_name"]), "center"),
            td(html_text(ip_label(config, row["source_ip"])), "center"),
            td(html_status(row["review_status"]), "center"),
            td(html_text(note)),
            td(html_text(row.get("remark", ""))),
        ]))
    return "".join(out)


def render_special_notes(summary_rows, config):
    notes = []
    if not authorized_user_state(config):
        notes.append("인가자 목록이 설정되지 않아 비인가자 접속 성공 여부는 자동 판정하지 않았습니다.")
    if summary_rows[1]["count"]:
        notes.append("업무시간 외 접속 이력은 사전 승인 또는 장애 대응 근거와 함께 별도 확인이 필요합니다.")
    if summary_rows[2]["count"]:
        notes.append("위험 패턴 명령은 명령 문자열만 추출하여 검토했으며, 세션 출력 내용은 보고서에 직접 포함하지 않았습니다.")
    if summary_rows[3]["count"]:
        notes.append("차단 이력은 IAM 권한, 세션 정책, 대상 인스턴스 상태와 함께 재검토하는 것이 좋습니다.")
    if not notes:
        notes.append("자동 수집 범위 내에서 즉시 조치가 필요한 이상 징후는 확인되지 않았습니다.")
    return html_list([html.escape(note) for note in notes])


def save_outputs(base_dir, success_rows, blocked_rows, command_rows, summary_rows):
    write_csv(
        success_rows,
        base_dir / "ssm_session_events.csv",
        [
            "event_time_kst", "event_time_display", "event_date_kst", "event_name", "user_name", "raw_user_name", "user_arn",
            "source_ip", "target_instance_id", "session_id", "mfa_authenticated",
            "aws_region", "after_hours", "weekend", "holiday", "source_ip_allowed", "user_authorized", "event_id",
        ],
    )
    write_csv(
        blocked_rows,
        base_dir / "ssm_blocked_events.csv",
        [
            "event_time_kst", "event_time_display", "event_date_kst", "event_name", "user_name", "raw_user_name", "user_arn",
            "source_ip", "target_instance_id", "session_id", "mfa_authenticated",
            "aws_region", "after_hours", "weekend", "holiday", "source_ip_allowed", "user_authorized", "event_id", "error_code", "error_message",
        ],
    )
    write_csv(
        command_rows,
        base_dir / "ssm_session_commands.csv",
        [
            "event_time", "event_time_kst", "session_id", "user_arn", "user_name", "run_as_user",
            "target_instance_id", "aws_region", "command", "risk_flags",
        ],
    )
    write_json({"summary_rows": summary_rows}, base_dir / "report_summary.json")


def autosize_sheet(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)


def append_sheet(ws, title, headers, rows):
    ws.title = title
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    autosize_sheet(ws)


def export_xlsx(base_dir, summary_rows, unauthorized_rows, after_hours_rows, risky_rows, blocked_rows, webapp_data):
    wb = Workbook()
    summary_ws = wb.active
    append_sheet(
        summary_ws,
        "Summary",
        ["항목", "건수", "검토결과", "비고"],
        [[row["name"], row["count"], row["status"], row["note"]] for row in summary_rows],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "WebAppSummary",
        ["서비스", "로그인 건수", "업무시간 외 로그인", "민감 처리행위", "검토결과", "검토 요약", "비고"],
        [
            [row["service"], row["login_count"], row["after_hours_count"], row["sensitive_count"], row["review_status"], row["review_summary"], row["remark"]]
            for row in (webapp_data or {}).get("service_summary", [])
        ] or [["-", "-", "-", "-", "-", "웹앱 로그 원본이 없어 검토 요약을 생성하지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "WebAppAfterHours",
        ["서비스", "탐지시간", "성명 / ID", "접속 IP", "검토결과", "검토 요약", "비고"],
        [
            [row["service"], row["event_time_display"], row["admin_label"], row["source_ip"], row["review_status"], row["review_summary"], row["remark"]]
            for row in (webapp_data or {}).get("after_hours_rows", [])
        ] or [["-", "-", "-", "-", "-", "업무시간 외 웹앱 로그인 이력이 확인되지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "WebAppSensitive",
        ["서비스", "탐지시간", "성명 / ID", "접속 IP", "검토결과", "검토 요약", "비고"],
        [
            [row["service"], row["event_time_display"], row["admin_label"], row["source_ip"], row["review_status"], row["review_summary"], row["remark"]]
            for row in (webapp_data or {}).get("sensitive_rows", [])
        ] or [["-", "-", "-", "-", "-", "민감 처리행위 로그가 확인되지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "Unauthorized",
        ["서비스", "탐지시간", "성명 / ID", "접속 IP", "검토결과", "검토 상세", "비고"],
        [
            ["PIMS bastion host", row["event_time_display"], row["user_name"], row["source_ip"], "비인가 의심", "인가자 목록에 없는 사용자로 확인되어 검토 필요", ""]
            for row in unauthorized_rows
        ] or [["-", "-", "-", "-", "-", "비인가자 접속 성공 이력이 확인되지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "AfterHours",
        ["서비스", "탐지시간", "성명 / ID", "접속 IP", "검토결과", "검토 상세", "비고"],
        [
            ["PIMS bastion host", row["event_time_display"], row["user_name"], row["source_ip"], row["review_status"], row["review_note"], row.get("remark", "")]
            for row in after_hours_rows
        ] or [["-", "-", "-", "-", "-", "업무시간 외 접속 이력이 확인되지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "RiskyCommands",
        ["서비스", "탐지시간", "성명 / ID", "세션 ID", "검토결과", "검토 상세", "비고"],
        [
            ["PIMS bastion host", row["event_time_kst"], row["user_arn"].rsplit("/", 1)[-1], row["session_id"], row["review_status"], f"{row['command']} / {row['review_note']}", row.get("remark", "")]
            for row in risky_rows
        ] or [["-", "-", "-", "-", "-", "위험 패턴으로 분류된 명령 이력이 확인되지 않았습니다.", ""]],
    )

    ws = wb.create_sheet()
    append_sheet(
        ws,
        "BlockedAttempts",
        ["서비스", "탐지시간", "성명 / ID", "접속 IP", "검토결과", "검토 상세", "비고"],
        [
            ["PIMS bastion host", row["event_time_display"], row["user_name"], row["source_ip"], row["review_status"], f"{row['error_code']}: {row.get('error_message', '')}".strip(), row.get("remark", "")]
            for row in blocked_rows
        ] or [["-", "-", "-", "-", "-", "차단된 StartSession 이력이 확인되지 않았습니다.", ""]],
    )

    xlsx_path = base_dir / "audit_report_details.xlsx"
    wb.save(xlsx_path)
    return xlsx_path


def main():
    parser = argparse.ArgumentParser(description="Generate bastion host access audit report HTML.")
    parser.add_argument("--profile", required=True)
    parser.add_argument("--bucket", required=True)
    parser.add_argument("--account-id", required=True)
    parser.add_argument("--region", required=True, help="CloudTrail region, e.g. ap-northeast-2")
    parser.add_argument("--log-group", default="/aws/ssm/bastion-audit-logs")
    parser.add_argument("--date-from", required=True)
    parser.add_argument("--date-to", required=True)
    parser.add_argument("--business-start-hour", type=int, default=9)
    parser.add_argument("--business-end-hour", type=int, default=19)
    parser.add_argument("--webapp-workbook", default="")
    parser.add_argument("--config", default="code/reviewer_config.json")
    parser.add_argument("--template", default="code/report_template.html")
    parser.add_argument("--out-dir", default="data/reports")
    args = parser.parse_args()

    start_date = dt.date.fromisoformat(args.date_from)
    end_date = dt.date.fromisoformat(args.date_to)
    config = load_config(args.config)

    report_key = f"{start_date:%Y%m}"
    base_dir = Path(args.out_dir) / report_key
    base_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = base_dir / "_cloudtrail_cache"

    all_records = []
    for day in daterange(start_date, end_date):
        day_dir = cache_dir / day.isoformat()
        sync_day_logs(args.bucket, args.account_id, args.region, day, args.profile, day_dir)
        for local_file in sorted(day_dir.glob("*.json.gz")):
            all_records.extend(load_cloudtrail_records(local_file))
        shutil.rmtree(day_dir)
    if cache_dir.exists():
        shutil.rmtree(cache_dir)

    success_rows, blocked_rows = extract_cloudtrail_findings(
        all_records,
        args.business_start_hour,
        args.business_end_hour,
        config,
    )
    command_rows = extract_command_rows(args.profile, args.log_group, success_rows)
    summary_rows, unauthorized_rows, after_hours_rows, risky_rows, blocked_detail_rows = summarize(
        success_rows,
        blocked_rows,
        command_rows,
        config,
    )
    webapp_data = extract_webapp_findings(
        args.webapp_workbook,
        start_date,
        end_date,
        args.business_start_hour,
        args.business_end_hour,
        config,
    )

    save_outputs(base_dir, success_rows, blocked_rows, command_rows, summary_rows)
    xlsx_path = export_xlsx(
        base_dir,
        summary_rows,
        unauthorized_rows,
        after_hours_rows,
        risky_rows,
        blocked_detail_rows,
        webapp_data,
    )

    template = load_template(args.template)
    generated_at = dt.datetime.now(tz=SEOUL).strftime("%Y-%m-%d %H:%M:%S %Z")
    report_title = f"{start_date.year}년 {start_date.month:02d}월 개인정보처리시스템 접속 로그 검토 보고"

    html_body = template.safe_substitute({
        "report_title": report_title,
        "review_period": f"- {start_date.year}년 {start_date.month:02d}월 {start_date.day}일 ~ {end_date.year}년 {end_date.month:02d}월 {end_date.day}일",
        "generated_at": f"자동생성 시각: {generated_at}",
        "scope_html": html_list([
            "개인정보처리시스템 접근 경유지: PIMS bastion host",
            "대상 로그 범위: SSM StartSession/ResumeSession/TerminateSession, Session Manager 세션 로그",
        ]),
        "scope_note": "운영툴 웹 로그인 로그는 본 자동 보고서 범위에 포함하지 않음",
        "review_items_html": html_list([
            "비인가자 접속 성공 유무",
            "업무시간 외 접속이력",
            "과도한 명령어 수행",
            "불필요 접속시도 차단 현황",
        ]),
        "review_item_note": "bastion host 자동 수집 기준",
        "criteria_html": html_list([
            f"인가자 접속 성공: reviewer_config.json의 authorized_users 기준",
            f"업무시간 외 접속: 평일 {args.business_start_hour:02d}:00 ~ {args.business_end_hour:02d}:00 외 시간 및 주말 접속",
            "과도한 명령어 수행: DB 접근, 권한 변경, 서비스 제어, 데이터 아카이브, 파괴적 명령 패턴 탐지",
            "불필요 접근시도 차단: CloudTrail의 StartSession 실패/거부 이벤트 기준",
        ]),
        "criteria_note": "인가자 기준 파일 미설정 시 비인가자 판정 생략",
        "log_sources_html": html_list([
            f"CloudTrail S3: {args.bucket} / {args.region}",
            f"CloudWatch Logs: {args.log_group}",
            f"웹앱 감사로그 원본: {Path(args.webapp_workbook).name}" if args.webapp_workbook else "웹앱 감사로그 원본: 미지정",
        ]),
        "log_source_note": "세부 근거와 원시성 상세 내역은 첨부 XLSX 기준으로 확인",
        "summary_rows": render_summary_rows(summary_rows),
        "webapp_summary_rows": render_webapp_summary_rows((webapp_data or {}).get("service_summary", [])),
        "webapp_unauthorized_rows": render_webapp_unauthorized_rows((webapp_data or {}).get("unauthorized_webapp_rows", []), bool(config.get("webapp_authorized_users"))),
        "webapp_after_hours_rows": render_webapp_event_rows((webapp_data or {}).get("after_hours_rows", []), "업무시간 외 웹앱 로그인 이력이 확인되지 않았습니다."),
        "webapp_sensitive_rows": render_webapp_event_rows(
            ((webapp_data or {}).get("sensitive_rows", []) + (webapp_data or {}).get("role_change_rows", []))[:30],
            "민감 처리행위 또는 권한 변경 로그가 확인되지 않았습니다.",
        ),
        "webapp_limit_rows": render_webapp_limit_rows((webapp_data or {}).get("limit_rows", [])),
        "special_notes_html": render_special_notes(summary_rows, config),
        "unauthorized_rows": render_unauthorized_rows(unauthorized_rows, config),
        "after_hours_rows": render_after_hours_rows(after_hours_rows, config),
        "risky_rows": render_risky_summary_rows(risky_rows),
        "blocked_rows": render_blocked_rows(blocked_detail_rows, config),
        "footer_note": "본 HTML은 검토 요약본이며, 상세 근거와 세부 내역은 첨부 XLSX를 기준으로 확인합니다. 인가자 목록과 예외 승인 정보는 reviewer_config.json에 의해 보강됩니다.",
    })

    report_path = base_dir / "audit_report.html"
    report_path.write_text(html_body, encoding="utf-8")

    print(json.dumps({
        "report_html": str(report_path),
        "details_xlsx": str(xlsx_path),
        "events_csv": str(base_dir / "ssm_session_events.csv"),
        "blocked_csv": str(base_dir / "ssm_blocked_events.csv"),
        "commands_csv": str(base_dir / "ssm_session_commands.csv"),
    }, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
